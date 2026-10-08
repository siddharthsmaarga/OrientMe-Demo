"""Reads whatever's actually in a topic's local folders and extracts plain
text from each file - the MVP ingestion model from the 17 Sep sprint
discussion: no SharePoint/Graph API connection, just files someone has
already downloaded or synced locally, read from an explicitly approved path.

Deliberately shallow (top-level + one level of subfolders, a small set of
text-bearing formats) - this is the backend-first MVP slice, not a full
document-management crawler.
"""

import email
import hashlib
import logging
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from email.policy import default as _email_default_policy
from html.parser import HTMLParser

from django.conf import settings as django_settings

from . import doc_meta
from .models import IngestedFile, Settings

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".eml"}  # .eml needs no extra package - stdlib email module

try:
    import extract_msg

    SUPPORTED_EXTENSIONS.add(".msg")
except ImportError:
    extract_msg = None

try:
    import docx  # python-docx

    SUPPORTED_EXTENSIONS.add(".docx")
except ImportError:
    docx = None

try:
    from pypdf import PdfReader

    SUPPORTED_EXTENSIONS.add(".pdf")
except ImportError:
    PdfReader = None

try:
    import openpyxl

    SUPPORTED_EXTENSIONS.update({".xlsx", ".xlsm"})
except ImportError:
    openpyxl = None

try:
    import pptx  # python-pptx

    SUPPORTED_EXTENSIONS.add(".pptx")
except ImportError:
    pptx = None

# Audio/video meeting recordings (e.g. Teams cloud recordings, a phone's own
# voice-memo recording of a call, or anything dropped in by a local
# transcription tool). Needs both faster-whisper (pip) and a real ffmpeg
# binary on PATH - gated on both, same fail-open-if-missing convention as
# every extractor above, so a machine without either dependency just
# doesn't pick up these extensions rather than crashing on import.
#
# The phone-recording formats (.aac/.amr/.3gp/.ogg/.flac) were added
# alongside Teams' own .mp4/.wav/.mp3/.m4a after Venki noted (29 Sep) that
# some meetings only exist as a phone recording, not a Teams cloud
# recording - ffmpeg decodes all of these through the exact same
# _extract_audio_video path below, so no separate code was needed, only a
# wider extension set.
AUDIO_VIDEO_EXTENSIONS = {".mp4", ".mov", ".webm", ".wav", ".mp3", ".m4a", ".aac", ".amr", ".3gp", ".ogg", ".flac"}
try:
    from faster_whisper import WhisperModel

    if shutil.which("ffmpeg"):
        SUPPORTED_EXTENSIONS.update(AUDIO_VIDEO_EXTENSIONS)
    else:
        WhisperModel = None
except ImportError:
    WhisperModel = None

_whisper_model = None  # loaded lazily, once per process - model load is the slow part


def _get_whisper_model():
    global _whisper_model
    if _whisper_model is None:
        model_size = getattr(django_settings, "WHISPER_MODEL_SIZE", "base")
        _whisper_model = WhisperModel(model_size, device="cpu", compute_type="int8")
    return _whisper_model


# Matches how Teams itself names an auto-saved recording/transcript, e.g.
# "Spend Potential Identification-20260928_093014UTC-Meeting Recording.mp4"
# - lets us recover a real recording date/time without needing diarization
# or any manual entry, falling back to the file's own mtime otherwise.
_TEAMS_FILENAME_RE = re.compile(r"(?P<title>.+?)-(?P<date>\d{8})_(?P<time>\d{6})UTC")


def _guess_recording_title_and_date(path):
    name = os.path.splitext(os.path.basename(path))[0]
    m = _TEAMS_FILENAME_RE.match(name)
    if m:
        title = m.group("title").strip()
        try:
            dt = datetime.strptime(m.group("date") + m.group("time"), "%Y%m%d%H%M%S")
            return title, dt
        except ValueError:
            pass
    mtime = datetime.fromtimestamp(os.path.getmtime(path))
    return name, mtime


def _format_timestamp(seconds):
    m, s = divmod(int(seconds), 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _extract_via_external_engine(path, url):
    """Hands the recording off to an external transcription/diarization
    service instead of the local faster-whisper model - the pluggable point
    for "our engineering team is also making a transcript diarizer tool
    that should also be integrated, or other tools" (Ayush, 29 Sep). Rather
    than hardcoding any one tool's internals (which don't exist yet), this
    is a plain HTTP contract anyone can point at: POST the audio file as
    multipart form data to whatever URL is configured in Settings
    (transcription_engine="external", external_transcription_url), and
    expect back JSON {"text": "...", "speakers": ["Name", ...]?} -
    "speakers" is optional and, when a real diarizer tool provides it,
    replaces the local engine's honest "not captured" disclaimer with an
    actual attendee list - the one thing local Whisper genuinely can't do.

    Raises on any failure (bad URL, timeout, malformed response) - the
    caller (extract_text) catches this and falls back to the local engine,
    so a misconfigured or temporarily-down external tool never blocks
    ingestion entirely."""
    import httpx

    title, when = _guess_recording_title_and_date(path)
    with open(path, "rb") as fh:
        response = httpx.post(url, files={"file": (os.path.basename(path), fh)}, timeout=600)
    response.raise_for_status()
    data = response.json()
    text = data.get("text", "")
    speakers = data.get("speakers")
    attendees = "; ".join(s for s in speakers if s) if speakers else "Not captured (engine did not report speakers)"
    header = (
        f"Title: {title}\n"
        f"Date: {when.strftime('%Y-%m-%d')}\n"
        f"Time: {when.strftime('%H:%M')}\n"
        f"Attendees: {attendees}\n"
    )
    return header + "\n" + text


def _extract_audio_video(path):
    """Transcribes a recording locally (faster-whisper, no cloud, no API
    key) and formats the output with a small synthetic header so it passes
    the app's existing looks_like_meeting_transcript() gate (extraction.py)
    the same way any other real transcript does - no special-casing needed
    downstream, it just flows through the one existing meeting-extraction
    and brief-generation pipeline every other ingested file already uses.

    Honest about what this can't do: faster-whisper has no speaker
    diarization, so there's no real "Attendees" list to report - the header
    says so plainly rather than inventing names."""
    title, when = _guess_recording_title_and_date(path)

    # Normalise to 16kHz mono WAV first - the format Whisper models expect,
    # and it means ffmpeg (not faster-whisper's own decoding) is the one
    # thing doing container/codec handling for every input type (.mp4,
    # .mov, .webm, .wav, .mp3, .m4a) via one code path.
    with tempfile.TemporaryDirectory() as tmp_dir:
        wav_path = os.path.join(tmp_dir, "audio.wav")
        proc = subprocess.run(
            ["ffmpeg", "-y", "-i", path, "-ac", "1", "-ar", "16000", wav_path],
            capture_output=True,
            text=True,
        )
        if proc.returncode != 0 or not os.path.exists(wav_path):
            raise ValueError(f"ffmpeg could not extract audio from {path}: {proc.stderr[-500:]}")

        model = _get_whisper_model()
        segments, _info = model.transcribe(wav_path, vad_filter=True)
        lines = [f"[{_format_timestamp(seg.start)}] {seg.text.strip()}" for seg in segments]

    header = (
        f"Title: {title}\n"
        f"Date: {when.strftime('%Y-%m-%d')}\n"
        f"Time: {when.strftime('%H:%M')}\n"
        f"Attendees: Not captured (auto-transcribed; no speaker diarization)\n"
    )
    return header + "\n" + "\n".join(lines)


def _extract_xlsx(path):
    # Every sheet, every row, tab-separated - keeps numbers/dates readable
    # as plain text for the same downstream extraction (meeting metadata,
    # employee hours, action items) that already runs on any other file,
    # rather than a separate spreadsheet-specific pipeline.
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cell.strip() for cell in cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


class _HtmlTextExtractor(HTMLParser):
    """Stdlib HTML -> plain text for an HTML-only email body: drops
    script/style/head content, turns block-level tags into line breaks."""

    _BLOCK_TAGS = {
        "p", "div", "br", "tr", "li", "ul", "ol", "table", "h1", "h2", "h3", "h4", "h5", "h6",
        "blockquote", "section", "article", "header", "footer", "hr", "pre",
    }
    _SKIP_TAGS = {"script", "style", "head", "title"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self._parts = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
        elif tag in self._BLOCK_TAGS:
            self._parts.append("\n")
        elif tag in ("td", "th"):
            self._parts.append("\t")

    def handle_endtag(self, tag):
        if tag in self._SKIP_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in self._BLOCK_TAGS:
            self._parts.append("\n")

    def handle_data(self, data):
        if self._skip_depth == 0:
            self._parts.append(data)

    def text(self):
        raw = "".join(self._parts).replace("\xa0", " ")
        lines = ["\t".join(" ".join(cell.split()) for cell in ln.split("\t")).strip("\t ") for ln in raw.splitlines()]
        out = []
        for ln in lines:
            if ln or (out and out[-1]):
                out.append(ln)
        return "\n".join(out).strip()


def html_to_text(html):
    """Strips tags from an HTML email body (an .eml/.msg with no plain-text
    part would otherwise ingest as tag soup)."""
    parser = _HtmlTextExtractor()
    try:
        parser.feed(html or "")
        parser.close()
    except Exception:  # noqa: BLE001 - malformed HTML: fall back to a crude strip
        return re.sub(r"<[^>]+>", " ", html or "").strip()
    return parser.text()


def _extract_eml(path):
    # Plain stdlib parsing, no LLM/heuristics involved - headers give the
    # who/when/subject signals the same way a transcript's own header does,
    # body text is walked to prefer the plain-text part over HTML (an
    # HTML-only mail is tag-stripped to text, see html_to_text).
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=_email_default_policy)
    header = (
        f"From: {msg.get('From', '')}\nTo: {msg.get('To', '')}\n"
        f"Date: {msg.get('Date', '')}\nSubject: {msg.get('Subject', '')}\n"
    )
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = body_part.get_content() if body_part else ""
    if body_part is not None and body_part.get_content_type() == "text/html":
        body = html_to_text(body)
    return header + "\n" + body


def _extract_msg(path):
    m = extract_msg.Message(path)
    try:
        header = (
            f"From: {m.sender or ''}\nTo: {m.to or ''}\n"
            f"Date: {m.date or ''}\nSubject: {m.subject or ''}\n"
        )
        body = m.body or ""
        if not body.strip():
            html = getattr(m, "htmlBody", None)
            if isinstance(html, bytes):
                html = html.decode("utf-8", errors="ignore")
            if html:
                body = html_to_text(html)
        return header + "\n" + body
    finally:
        m.close()


def _extract_docx(path):
    """Paragraph text plus table cells, in document order - quotations/SOWs
    are mostly tables. A paragraph-only document produces exactly the same
    text as before (non-empty paragraphs joined by newlines); a table row
    becomes one tab-separated line (merged cells are not repeated)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    lines = []

    def add_table(table):
        for row in table.rows:
            cells = []
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                cells.append(" ".join(cell.text.split()))
            if any(cells):
                lines.append("\t".join(cells))

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, d).text
            if text.strip():
                lines.append(text)
        elif tag == "tbl":
            add_table(Table(child, d))
    return "\n".join(lines)


def _extract_pptx(path):
    prs = pptx.Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append("\t".join(cell.text for cell in row.cells))
        if lines:
            parts.append(f"--- Slide {i} ---\n" + "\n".join(lines))
    return "\n\n".join(parts)


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".txt", ".md"):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    if ext == ".eml":
        return _extract_eml(path)
    if ext == ".msg" and extract_msg is not None:
        return _extract_msg(path)
    if ext == ".docx" and docx is not None:
        return _extract_docx(path)
    if ext == ".pdf" and PdfReader is not None:
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext in (".xlsx", ".xlsm") and openpyxl is not None:
        return _extract_xlsx(path)
    if ext == ".pptx" and pptx is not None:
        return _extract_pptx(path)
    if ext in AUDIO_VIDEO_EXTENSIONS and WhisperModel is not None:
        # Engine selection (Settings page / Transcripts page): "external"
        # tries whatever URL is configured first (the future diarizer tool,
        # or anything else standing behind the same HTTP contract), falling
        # back to the local model on any failure so a misconfigured or
        # down external tool never blocks ingestion - never a silent swap,
        # logged so it's visible in backend.log.
        engine = Settings.get("transcription_engine", "local_whisper")
        if engine == "external":
            external_url = Settings.get("external_transcription_url", "").strip()
            if external_url:
                try:
                    return _extract_via_external_engine(path, external_url)
                except Exception:
                    logger.exception(
                        "External transcription engine failed for %s - falling back to local Whisper", path
                    )
        return _extract_audio_video(path)
    raise ValueError(f"Unsupported or unavailable extractor for {ext}")


def _walk_files(folder_path, max_depth=1):
    """Top level plus one level of subfolders - enough for a real project
    folder without accidentally ingesting an entire OneDrive root."""
    base_depth = folder_path.rstrip(os.sep).count(os.sep)
    for root, dirs, files in os.walk(folder_path):
        depth = root.rstrip(os.sep).count(os.sep) - base_depth
        if depth >= max_depth:
            dirs[:] = []
        for name in files:
            if os.path.splitext(name)[1].lower() in SUPPORTED_EXTENSIONS:
                yield os.path.join(root, name)


def ingest_folder(topic, folder_path, on_progress=None):
    """Reads every supported file under folder_path, skipping ones whose
    content hash hasn't changed since last time - so re-ingesting a large
    folder after adding one new file is cheap, matching the PRD's own
    'refresh only changed material' goal. on_progress(done, total), if
    given, is called after each file - this is the fast local-I/O phase, so
    it mostly matters for a real progress bar on a very large folder."""
    results = {"scanned": 0, "added": 0, "updated": 0, "unchanged": 0, "errors": []}

    if not os.path.isdir(folder_path):
        results["errors"].append(f"Folder not found: {folder_path}")
        return results

    paths = list(_walk_files(folder_path))
    total = len(paths)
    for i, path in enumerate(paths, 1):
        results["scanned"] += 1
        try:
            with open(path, "rb") as f:
                raw = f.read()
            content_hash = hashlib.sha256(raw).hexdigest()

            existing = IngestedFile.objects.filter(topic=topic, path=path).first()
            if existing and existing.content_hash == content_hash:
                results["unchanged"] += 1
                continue
            is_new = existing is None

            text = extract_text(path)
            obj, _ = IngestedFile.objects.update_or_create(
                topic=topic,
                path=path,
                defaults={
                    "extracted_text": text,
                    "content_hash": content_hash,
                    "extraction_error": "",
                },
            )
            doc_meta.safe_apply_auto_metadata(obj)
            results["added" if is_new else "updated"] += 1
        except Exception as e:  # noqa: BLE001 - a bad file shouldn't abort the whole ingest
            bad, _ = IngestedFile.objects.update_or_create(
                topic=topic,
                path=path,
                defaults={"extraction_error": str(e)},
            )
            doc_meta.safe_apply_auto_metadata(bad)
            results["errors"].append(f"{path}: {e}")
        finally:
            if on_progress:
                on_progress(i, total)

    return results
