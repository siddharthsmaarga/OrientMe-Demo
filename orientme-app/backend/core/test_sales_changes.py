"""Tests for the 5 Oct 2026 additive sales-oriented changes (build spec A-H) and
the conversation-style "who spoke what" transcript (section I).

No test calls a real LLM: the API model is replaced by a mock function.
"""

import csv
import io
import os
import tempfile
from datetime import date
from unittest import mock

from django.contrib.auth.models import User
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.test import TestCase, TransactionTestCase, override_settings
from rest_framework.test import APIClient

from . import doc_meta, extraction, ingestion, llm_brief, recordings
from .models import (
    CaseSummary,
    Contact,
    ExtractedMeta,
    IngestedFile,
    MeetingNote,
    RecordingEvent,
    Settings,
    Task,
    Topic,
)
from .views import _dedupe_meetings, _extract_and_store

TEAMS_TEXT = (
    "Orient-Me _ Shared Context Project Meeting-20260923_045358UTC-Meeting Recording\n"
    "September 23, 2026, 4:53AM\n7m 25s\n\n"
    "Siddharth Subramanian   0:29\n"
    "And this is exactly the shape that I want, right?\nI get that.\n"
    "It will show from where the data is fetched and you can edit that also.\n\n"
    "Ayush Sinha A   1:13\n"
    "Okay, so I will build the first version of the timeline by Friday for the demo.\n\n"
    "Okay. Vijay Krishna Ramachandran   10:16:03\n"
    "We should also keep the payroll workflow exactly as it is for the finance team.\n"
    "Ayush Sinha A stopped transcription\n"
)

WHISPER_TEXT = (
    "Title: Weekly Update\nDate: 2026-10-05\nTime: 04:45\n"
    "Attendees: Not captured (auto-transcribed; no speaker diarization)\n\n"
    + "\n".join(f"[00:{i:02d}] This is spoken sentence number {i} about the roadmap and the demo." for i in range(10, 40))
)


def make_user(name="tester"):
    return User.objects.create_user(name, password="pw-not-used")


def api_client(user=None):
    client = APIClient()
    client.force_authenticate(user or make_user())
    return client


def make_file(topic, name="file.txt", text="", **kw):
    return IngestedFile.objects.create(
        topic=topic, path=kw.pop("path", f"/tmp/{name}"), display_name=name, extracted_text=text, **kw
    )


# ---------------------------------------------------------------------------
# A. Classifier, versions, dates, signal, tags
# ---------------------------------------------------------------------------
class ClassifierTests(TestCase):
    def classify(self, text="", name="x.txt", **kw):
        return doc_meta.classify_file(text, [name], **kw)

    def test_email_by_header_extension_and_source(self):
        header = "From: Priya <p@x.com>\nTo: venki@maarga.com\nDate: Mon, 03 Aug 2026 10:00:00 +0530\nSubject: Intro call\n\nHi,"
        self.assertEqual(self.classify(header, "notes.txt")["doc_type"], "email")
        self.assertEqual(self.classify("", "thread.eml")["doc_type"], "email")
        self.assertEqual(self.classify("", "thread.msg")["doc_type"], "email")
        self.assertEqual(self.classify("body", "Re something", source_method="email_import")["doc_type"], "email")

    def test_transcript_signals(self):
        self.assertEqual(self.classify(TEAMS_TEXT, "meeting.docx")["doc_type"], "meeting_transcript")
        self.assertEqual(self.classify(WHISPER_TEXT, "rec.txt")["doc_type"], "meeting_transcript")
        # a single speaker line is still the app's own definition of a meeting transcript
        one = "Title line\nSiddharth Subramanian   0:29\nSome words said here today.\n"
        self.assertEqual(self.classify(one, "call.docx")["doc_type"], "meeting_transcript")

    def test_filename_keywords(self):
        cases = {
            "Northwind_Proposal_V1.docx": "proposal",
            "Northwind Quotation.xlsx": "quotation_sow",
            "Pilot_SOW_v2.pdf": "quotation_sow",
            "quote-final.pdf": "quotation_sow",
            "Northwind_BRD_V1.docx": "prd_brd",
            "product PRD.docx": "prd_brd",
            "Project_Charter.docx": "project_charter",
            "Northwind_Project_Plan.xlsx": "project_plan",
            "Intro_call_MoM.docx": "mom_action_items",
            "meeting minutes 5 Oct.docx": "mom_action_items",
            "Stakeholder Map.pptx": "stakeholder_map",
            "Kickoff.docx": "meeting_transcript",
            "weekly-standup.txt": "meeting_transcript",
            "Q3 report.pdf": "report",
            "random.bin": "other",
        }
        for name, expected in cases.items():
            self.assertEqual(self.classify("", name)["doc_type"], expected, name)

    def test_mom_filename_beats_attendees_header(self):
        text = "Meeting: Intro\nAttendees: A, B\nAction items:\n- send deck\n"
        self.assertEqual(self.classify(text, "Intro_call_MoM.docx")["doc_type"], "mom_action_items")

    def test_proposal_with_agenda_time_line_stays_proposal(self):
        text = "Northwind Proposal\nKickoff Meeting   9:30\nScope of the pilot."
        self.assertEqual(self.classify(text, "Northwind_Proposal_V1.docx")["doc_type"], "proposal")

    def test_content_rules_only_look_at_title_area(self):
        mention = "Gap analysis of tools\n\nSome intro text.\nLater we discuss the quotation process in detail."
        self.assertEqual(self.classify(mention, "analysis.md")["doc_type"], "notes")
        titled = "Statement of Work\nFixed fee for the pilot."
        self.assertEqual(self.classify(titled, "doc.docx")["doc_type"], "quotation_sow")

    def test_paste_and_text_fallbacks(self):
        self.assertEqual(self.classify("just words", "x", source_method="paste")["doc_type"], "notes")
        self.assertEqual(self.classify("just words", "readme.md")["doc_type"], "notes")

    def test_version_label(self):
        v = doc_meta.detect_version_label
        self.assertEqual(v(["Northwind_Proposal_V1.docx"]), "V1")
        self.assertEqual(v(["proposal v2.docx"]), "V2")
        self.assertEqual(v(["Proposal_Final.pdf"]), "Final")
        self.assertEqual(v(["quote_v3_final.xlsx"]), "Final")
        self.assertEqual(v(["plan_v10.docx"]), "V10")
        self.assertEqual(v(["dev1_notes.txt"]), "")
        self.assertEqual(v(["review.docx"]), "")
        self.assertEqual(v(["Northwind_Quotation_SOW_v1.xlsx"]), "V1")

    def test_document_date(self):
        d = doc_meta.detect_document_date
        email_text = "From: a@b.com\nTo: c@d.com\nDate: Mon, 03 Aug 2026 10:00:00 +0530\nSubject: Hi\n\nbody"
        self.assertEqual(d(email_text, ["mail.eml"]), date(2026, 8, 3))
        self.assertEqual(d("Title\nDate: 2026-09-16\nbody", ["x.docx"]), date(2026, 9, 16))
        self.assertEqual(d("body", ["2026-08-14_Discovery_1.docx"]), date(2026, 8, 14))
        self.assertEqual(d("body", ["Meeting-20260923_045358UTC.docx"]), date(2026, 9, 23))
        self.assertEqual(d("rec-20261005_044517UTC-Meeting Recording\nSiddharth Subramanian 0:01", ["x.docx"]), date(2026, 10, 5))
        self.assertIsNone(d("no dates here", ["plain.txt"]))
        self.assertIsNone(d("Date: not a date", ["plain.txt"]))

    def test_signal_for_email_and_document(self):
        text = (
            "From: Priya <p@x.com>\nTo: v@m.com\nDate: Mon, 03 Aug 2026 10:00:00 +0530\n"
            "Subject: Intro call request\n\nHi Venki,\n\nWe would like a 30 minute intro call next week to "
            "discuss AI services.\n\n-----Original Message-----\nquoted stuff that must be ignored here today."
        )
        sig = doc_meta.build_signal(text, "email")
        self.assertTrue(sig.startswith("Intro call request - "))
        self.assertIn("30 minute intro call", sig)
        self.assertNotIn("quoted stuff", sig)
        long = "Word " * 100 + "end."
        self.assertLessEqual(len(doc_meta.build_signal(long, "other")), doc_meta.MAX_SIGNAL_CHARS)
        doc = "Northwind Proposal V1\n\nThis proposal describes a six week pilot for exception handling in the warehouse."
        self.assertIn("six week pilot", doc_meta.build_signal(doc, "proposal"))

    def test_signal_for_transcript_skips_header_lines(self):
        sig = doc_meta.build_signal(TEAMS_TEXT, "meeting_transcript")
        self.assertEqual(sig, "And this is exactly the shape that I want, right?")  # speech, not the title/date header

    def test_tags_max_three_from_vocabulary(self):
        text = "The price and the budget and the quotation and payment terms. Next steps: follow up. Timeline and deadline and milestone."
        tags = doc_meta.suggest_tags(text, "quotation_sow")
        self.assertLessEqual(len(tags), 3)
        self.assertIn("Commercial", tags)
        self.assertEqual(doc_meta.suggest_tags("hello world", "other"), [])
        self.assertEqual(doc_meta.normalize_tags("a, b, A, c, d"), ["a", "b", "c"])
        self.assertEqual(doc_meta.normalize_tags(["next steps", "RISKS"]), ["Next Steps", "Risks"])

    def test_apply_auto_metadata_sets_fields_state_and_seconds(self):
        topic = Topic.objects.create(name="T")
        f = make_file(topic, "Northwind_Proposal_V2.docx", "Northwind Proposal\n\nThe pilot costs a fixed fee and the price is final for the pilot budget.")
        self.assertTrue(doc_meta.apply_auto_metadata(f))
        f.refresh_from_db()
        self.assertEqual((f.doc_type, f.doc_type_source, f.version_label), ("proposal", "auto", "V2"))
        self.assertEqual(f.meta_state, "generated_unconfirmed")
        self.assertIsNotNone(f.meta_seconds)
        self.assertFalse(doc_meta.apply_auto_metadata(f))  # idempotent

    def test_error_file_gets_error_state(self):
        topic = Topic.objects.create(name="T")
        f = make_file(topic, "Northwind_Project_Plan.xlsx", "", extraction_error="Could not read the sheet")
        doc_meta.apply_auto_metadata(f)
        f.refresh_from_db()
        self.assertEqual((f.meta_state, f.doc_type), ("error", "project_plan"))

    def test_never_overwrites_manual_or_confirmed_values(self):
        topic = Topic.objects.create(name="T")
        f = make_file(topic, "Proposal_V1.docx", "Proposal\n\nA sentence that is long enough to be a signal here.")
        doc_meta.apply_auto_metadata(f)
        f.doc_type, f.doc_type_source = "report", "manual"
        f.extracted_text = "Statement of Work\n\nBrand new content that changes the signal sentence here."
        f.save()
        doc_meta.apply_auto_metadata(f)
        f.refresh_from_db()
        self.assertEqual(f.doc_type, "report")  # manual doc_type kept even in the generated state
        self.assertIn("Brand new content", f.signal)  # the other auto fields still follow the content
        f.signal, f.meta_state = "My own signal", "corrected"
        f.extracted_text = "Another completely different sentence that would otherwise become the signal."
        f.save()
        doc_meta.apply_auto_metadata(f)
        f.refresh_from_db()
        self.assertEqual(f.signal, "My own signal")  # corrected rows are left alone entirely

    def test_ingest_folder_fills_metadata_for_each_file(self):
        topic = Topic.objects.create(name="T")
        with tempfile.TemporaryDirectory() as d:
            with open(os.path.join(d, "Northwind_Quotation_v1.txt"), "w", encoding="utf-8") as fh:
                fh.write("Quotation\n\nFixed fee for the six week pilot, payable in two instalments of the price.")
            ingestion.ingest_folder(topic, d)
        f = topic.files.get()
        self.assertEqual((f.doc_type, f.version_label, f.meta_state), ("quotation_sow", "V1", "generated_unconfirmed"))

    def test_artifact_type_reports_email_for_eml(self):
        from .serializers import artifact_type_for

        topic = Topic.objects.create(name="T")
        eml = make_file(topic, "a.eml", "x", path="/tmp/a.eml", source_method="folder")
        plain = make_file(topic, "b.txt", "x", path="/tmp/b.txt")
        self.assertEqual(artifact_type_for(eml), "email")
        self.assertEqual(artifact_type_for(plain), "file")  # old derivation unchanged
        ExtractedMeta.objects.create(ingested_file=plain)
        plain.refresh_from_db()
        self.assertEqual(artifact_type_for(plain), "meeting")


# ---------------------------------------------------------------------------
# B. Meeting type / sequence + dedupe
# ---------------------------------------------------------------------------
class MeetingTypeTests(TestCase):
    def test_detect_meeting_type_and_sequence(self):
        d = doc_meta.detect_meeting_type
        self.assertEqual(d("Northwind Discovery 2"), ("discovery", 2))
        self.assertEqual(d("Discovery_3 call"), ("discovery", 3))
        self.assertEqual(d("Discovery call 1"), ("discovery", 1))
        self.assertEqual(d("Solution review 2"), ("solution", 2))
        self.assertEqual(d("Initial intro call"), ("intro", None))
        self.assertEqual(d("Project kick-off"), ("kickoff", None))
        self.assertEqual(d("Weekly stand-up"), ("standup", None))
        self.assertEqual(d("Customer sync"), ("standup", None))
        self.assertEqual(d("Pricing discussion"), ("commercial", None))
        self.assertEqual(d("Design review"), ("review", None))
        self.assertEqual(d("2nd discovery session"), ("discovery", 2))
        self.assertEqual(d("AI Adoption Sprint Meeting (4)"), ("", None))  # a download counter is not a sequence
        self.assertEqual(d("Orient-Me Check-in"), ("", None))
        self.assertEqual(d("", "Discovery_2.docx"), ("discovery", 2))
        self.assertEqual(d("Plain title", "Kickoff"), ("kickoff", None))

    def test_extract_and_store_sets_auto_type_and_keeps_manual(self):
        topic = Topic.objects.create(name="T")
        text = "Meeting: Discovery 2 with Northwind\nDate: 2026-08-21\n" + TEAMS_TEXT.split("\n\n", 1)[1]
        f = make_file(topic, "d2.txt", text)
        _extract_and_store(f)
        m = f.meta
        self.assertEqual((m.meeting_type, m.meeting_sequence, m.meeting_type_source), ("discovery", 2, "auto"))
        self.assertGreaterEqual(len(m.speaker_turns), 3)
        m.meeting_type, m.meeting_sequence, m.meeting_type_source = "solution", 1, "manual"
        m.save()
        _extract_and_store(f)
        m.refresh_from_db()
        self.assertEqual((m.meeting_type, m.meeting_sequence, m.meeting_type_source), ("solution", 1, "manual"))

    def test_dedupe_keeps_old_behaviour_for_untyped_titles(self):
        topic = Topic.objects.create(name="T")
        a = ExtractedMeta.objects.create(
            ingested_file=make_file(topic, "a"), meeting_title="Weekly Orient Check In Meeting",
            meeting_date=date(2026, 10, 1), attendees=["A"],
        )
        b = ExtractedMeta.objects.create(
            ingested_file=make_file(topic, "b"), meeting_title="weekly orient check in Meeting",
            meeting_date=date(2026, 10, 1), attendees=["A", "B"],
        )
        kept = _dedupe_meetings([a, b])
        self.assertEqual(len(kept), 1)
        self.assertEqual(kept[0].id, b.id)

    def test_dedupe_does_not_merge_typed_meetings_with_different_sequence_or_type(self):
        topic = Topic.objects.create(name="T")
        common = dict(meeting_date=date(2026, 8, 21))
        d1 = ExtractedMeta.objects.create(
            ingested_file=make_file(topic, "d1"), meeting_title="Acme Corp Discovery Call 1",
            meeting_type="discovery", meeting_sequence=1, **common)
        d2 = ExtractedMeta.objects.create(
            ingested_file=make_file(topic, "d2"), meeting_title="Acme Corp Discovery Call 2",
            meeting_type="discovery", meeting_sequence=2, **common)
        s1 = ExtractedMeta.objects.create(
            ingested_file=make_file(topic, "s1"), meeting_title="Acme Corp Solution Call",
            meeting_type="solution", **common)
        self.assertEqual(len(_dedupe_meetings([d1, d2, s1])), 3)

    def test_meeting_display_title(self):
        t = doc_meta.meeting_display_title
        self.assertEqual(t("Acme Corp Discovery", "Acme Corp Discovery", "", None), "Acme Corp Discovery")
        self.assertEqual(t("Acme Corp Call", "Acme Corp Call", "discovery", 2), "Discovery 2: Acme Corp Call")
        self.assertEqual(t("Discovery 2 call", "Discovery 2 call with Acme", "discovery", 2), "Discovery 2 call with Acme")


class MeetingEndpointTests(TestCase):
    def setUp(self):
        self.client = api_client()
        self.topic = Topic.objects.create(name="Northwind", topic_type="opportunity")
        f = make_file(self.topic, "d.txt", TEAMS_TEXT)
        self.meta = ExtractedMeta.objects.create(
            ingested_file=f, meeting_title="Discovery call", meeting_date=date(2026, 8, 21),
            attendees=["Not captured (auto-transcribed; no speaker diarization)", "Priya"],
            speaker_turns=extraction.parse_speaker_turns(TEAMS_TEXT),
        )

    def test_meeting_metadata_old_keys_new_keys_and_clean_attendees(self):
        res = self.client.get(f"/api/topics/{self.topic.id}/meeting_metadata/")
        self.assertEqual(res.status_code, 200)
        row = res.json()[0]
        for key in ("id", "file_display_name", "file_path", "meeting_title", "timeline_title", "meeting_date",
                    "meeting_time", "attendees", "agenda", "achieved"):
            self.assertIn(key, row)
        for key in ("file_id", "meeting_type", "meeting_type_label", "meeting_sequence", "meeting_type_source",
                    "speaker_turn_count", "speaker_summary", "notes"):
            self.assertIn(key, row)
        self.assertEqual(row["attendees"], ["Priya"])  # placeholder hidden on read
        self.assertNotIn("speaker_turns", row)
        self.assertEqual(row["speaker_turn_count"], 3)
        self.assertEqual(row["speaker_summary"][0]["speaker"], "Siddharth Subramanian")
        full = self.client.get(f"/api/topics/{self.topic.id}/meeting_metadata/?include_turns=1").json()[0]
        self.assertEqual(len(full["speaker_turns"]), 3)

    def test_patch_meeting_type_sets_manual(self):
        res = self.client.patch(f"/api/meetings/{self.meta.id}/", {"meeting_type": "discovery", "meeting_sequence": 3}, format="json")
        self.assertEqual(res.status_code, 200)
        body = res.json()
        self.assertEqual((body["meeting_type"], body["meeting_sequence"], body["meeting_type_source"]), ("discovery", 3, "manual"))
        self.assertEqual(body["timeline_title"], "Discovery 3: Discovery call" if "Discovery 3" not in body["meeting_title"] else body["meeting_title"])
        bad = self.client.patch(f"/api/meetings/{self.meta.id}/", {"meeting_type": "nonsense"}, format="json")
        self.assertEqual(bad.status_code, 400)
        bad = self.client.patch(f"/api/meetings/{self.meta.id}/", {"meeting_sequence": 0}, format="json")
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(self.client.patch(f"/api/meetings/{self.meta.id}/", {}, format="json").status_code, 400)
        cleared = self.client.patch(f"/api/meetings/{self.meta.id}/", {"meeting_type": "", "meeting_sequence": None}, format="json")
        self.assertEqual(cleared.json()["meeting_type"], "")


# ---------------------------------------------------------------------------
# C. Topic stage + customer_name
# ---------------------------------------------------------------------------
class TopicStageTests(TestCase):
    def test_stage_and_customer_name_roundtrip_and_old_fields_remain(self):
        client = api_client()
        res = client.post("/api/topics/", {"name": "Northwind", "topic_type": "opportunity", "stage": "discovery", "customer_name": "Northwind Retail"}, format="json")
        self.assertEqual(res.status_code, 201)
        body = res.json()
        for key in ("id", "name", "topic_type", "one_liner", "related_people", "created_at", "auto_created",
                    "folders", "file_count", "last_activity_at", "latest_risks_and_gaps"):
            self.assertIn(key, body)
        self.assertEqual((body["stage"], body["customer_name"]), ("discovery", "Northwind Retail"))
        patched = client.patch(f"/api/topics/{body['id']}/", {"stage": "won"}, format="json").json()
        self.assertEqual(patched["stage"], "won")
        self.assertEqual(client.patch(f"/api/topics/{body['id']}/", {"stage": "bogus"}, format="json").status_code, 400)
        old = client.post("/api/topics/", {"name": "Old style"}, format="json").json()
        self.assertEqual((old["stage"], old["customer_name"], old["topic_type"]), ("", "", "project"))


# ---------------------------------------------------------------------------
# D. MeetingNote CRUD
# ---------------------------------------------------------------------------
class MeetingNoteTests(TestCase):
    def setUp(self):
        self.user = make_user("ayush")
        self.client = api_client(self.user)
        self.topic = Topic.objects.create(name="Northwind", topic_type="opportunity")
        self.other = Topic.objects.create(name="Other")
        self.meta = ExtractedMeta.objects.create(
            ingested_file=make_file(self.topic, "d.txt", "x"), meeting_title="Discovery 1", meeting_date=date(2026, 8, 14))
        self.other_meta = ExtractedMeta.objects.create(ingested_file=make_file(self.other, "o.txt", "x"))

    def test_crud(self):
        url = f"/api/topics/{self.topic.id}/notes/"
        created = self.client.post(url, {"text": "  Sponsor is the ops head.  ", "meta": self.meta.id}, format="json")
        self.assertEqual(created.status_code, 201)
        note = created.json()
        self.assertEqual((note["text"], note["author"], note["meta"]), ("Sponsor is the ops head.", "ayush", self.meta.id))
        self.assertEqual(note["meeting_title"], "Discovery 1")
        self.assertEqual(note["meeting_date"], "2026-08-14")
        free = self.client.post(url, {"text": "General note"}, format="json").json()
        self.assertIsNone(free["meta"])
        self.assertEqual(len(self.client.get(url).json()), 2)
        self.assertEqual(len(self.client.get(url + f"?meta={self.meta.id}").json()), 1)
        patched = self.client.patch(f"{url}{note['id']}/", {"text": "Updated"}, format="json")
        self.assertEqual(patched.json()["text"], "Updated")
        self.assertEqual(patched.json()["author"], "ayush")
        self.assertEqual(self.client.delete(f"{url}{note['id']}/").status_code, 204)
        self.assertEqual(len(self.client.get(url).json()), 1)

    def test_validation(self):
        url = f"/api/topics/{self.topic.id}/notes/"
        self.assertEqual(self.client.post(url, {"text": "   "}, format="json").status_code, 400)
        self.assertEqual(self.client.post(url, {"text": "x" * 2001}, format="json").status_code, 400)
        self.assertEqual(self.client.post(url, {"text": "ok", "meta": self.other_meta.id}, format="json").status_code, 400)
        note = MeetingNote.objects.create(topic=self.other, text="not yours")
        self.assertEqual(self.client.patch(f"{url}{note.id}/", {"text": "x"}, format="json").status_code, 404)
        self.assertEqual(self.client.delete(f"{url}{note.id}/").status_code, 404)

    def test_note_survives_meeting_row_removal(self):
        note = MeetingNote.objects.create(topic=self.topic, meta=self.meta, text="keep me")
        self.meta.delete()
        note.refresh_from_db()
        self.assertIsNone(note.meta_id)

    def test_notes_in_meeting_metadata(self):
        MeetingNote.objects.create(topic=self.topic, meta=self.meta, text="after the call", author="a")
        row = self.client.get(f"/api/topics/{self.topic.id}/meeting_metadata/").json()[0]
        self.assertEqual([n["text"] for n in row["notes"]], ["after the call"])


# ---------------------------------------------------------------------------
# E. Timeline
# ---------------------------------------------------------------------------
class TimelineTests(TestCase):
    def setUp(self):
        self.client = api_client()
        self.topic = Topic.objects.create(name="Northwind", topic_type="opportunity")
        t = self.topic
        self.email = make_file(
            t, "intro.eml", "From: Priya <p@x.com>\nTo: v@m.com\nDate: Mon, 03 Aug 2026 10:00:00 +0530\nSubject: Intro call request\n\nHello there, we want an intro call next week.",
            doc_type="email", document_date=date(2026, 8, 3), signal="Intro call request - we want a call")
        self.prop = make_file(t, "Northwind_Proposal_V2.docx", "x", doc_type="proposal", version_label="V2", document_date=date(2026, 9, 9), signal="V2 moves pricing to fixed fee")
        make_file(t, "plan.xlsx", "x", doc_type="project_plan", document_date=date(2026, 10, 3))
        make_file(t, "notes.md", "x", doc_type="notes", document_date=date(2026, 8, 1))  # not a timeline kind
        make_file(t, "broken.docx", "", doc_type="proposal", meta_state="error", document_date=date(2026, 8, 2))
        self.undated = make_file(t, "quote.pdf", "x", doc_type="quotation_sow")
        self.meta = ExtractedMeta.objects.create(
            ingested_file=make_file(t, "d.txt", "x"), meeting_title="Acme Discovery Call", meeting_date=date(2026, 8, 14),
            meeting_type="discovery", meeting_sequence=1, achieved="Pain points listed")
        MeetingNote.objects.create(topic=t, meta=self.meta, text="Sponsor is the ops head", author="ayush")

    def test_topic_timeline_events(self):
        res = self.client.get(f"/api/topics/{self.topic.id}/timeline_events/")
        self.assertEqual(res.status_code, 200)
        events = res.json()
        kinds = [e["kind"] for e in events]
        self.assertEqual(set(kinds), {"meeting", "email", "document", "note"})
        dated = [e["date"] for e in events]
        self.assertEqual(dated, sorted(dated))  # chronological ascending
        meeting = next(e for e in events if e["kind"] == "meeting")
        self.assertEqual(meeting["title"], "Discovery 1: Acme Discovery Call")
        self.assertEqual((meeting["meeting_type"], meeting["meeting_sequence"]), ("discovery", 1))
        email = next(e for e in events if e["kind"] == "email")
        self.assertEqual((email["date"], email["title"], email["approx_date"]), ("2026-08-03", "Intro call request", False))
        proposal = next(e for e in events if e["doc_type"] == "proposal")
        self.assertEqual((proposal["title"], proposal["version_label"], proposal["date"]), ("Proposal V2", "V2", "2026-09-09"))
        undated = next(e for e in events if e["doc_type"] == "quotation_sow")
        self.assertTrue(undated["approx_date"])
        self.assertNotIn("notes", [e["doc_type"] for e in events])
        self.assertEqual(len([e for e in events if e["doc_type"] == "proposal"]), 1)  # error file skipped
        note = next(e for e in events if e["kind"] == "note")
        self.assertEqual((note["author"], note["meeting_title"]), ("ayush", "Acme Discovery Call"))

    def test_global_timeline_has_new_kinds_and_old_shape(self):
        events = self.client.get("/api/timeline/").json()
        self.assertTrue({"meeting", "email", "document", "note"} <= {e["kind"] for e in events})
        for e in events:
            for key in ("kind", "date", "topic_id", "topic_name", "title"):
                self.assertIn(key, e)
        self.assertIn("Discovery 1: Acme Discovery Call", [e["title"] for e in events])

    def test_existing_meeting_without_type_keeps_title(self):
        t2 = Topic.objects.create(name="Plain")
        ExtractedMeta.objects.create(ingested_file=make_file(t2, "p", "x"), meeting_title="AI Adoption sprint meeting 4", meeting_date=date(2026, 9, 30))
        row = self.client.get(f"/api/topics/{t2.id}/meeting_metadata/").json()[0]
        self.assertEqual(row["timeline_title"], "AI Adoption Sprint")  # the old hard-coded label is untouched


# ---------------------------------------------------------------------------
# A. API: file metadata edit / confirm / CSV / sources
# ---------------------------------------------------------------------------
class FileMetadataApiTests(TestCase):
    def setUp(self):
        self.client = api_client()
        self.topic = Topic.objects.create(name="Northwind", topic_type="opportunity")
        self.f = make_file(self.topic, "Northwind_Proposal_V1.docx", "Proposal\n\nThe pilot costs a fixed fee for the price.")
        doc_meta.apply_auto_metadata(self.f)

    def test_patch_sets_corrected_and_manual(self):
        res = self.client.patch(
            f"/api/files/{self.f.id}/metadata/",
            {"doc_type": "quotation_sow", "topic_tags": ["commercial", "Timeline", "Risks", "Extra"], "signal": "My line",
             "version_label": "Final", "document_date": "2026-09-16"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual((body["doc_type"], body["doc_type_source"], body["meta_state"]), ("quotation_sow", "manual", "corrected"))
        self.assertEqual(body["topic_tags"], "Commercial, Timeline, Risks")
        self.assertEqual((body["version_label"], body["document_date"], body["doc_type_label"]), ("Final", "2026-09-16", "Quotation / SOW"))
        again = self.client.patch(f"/api/files/{self.f.id}/metadata/", {"document_date": None}, format="json").json()
        self.assertIsNone(again["document_date"])

    def test_review_notes_alone_keep_state(self):
        before = self.f.meta_state
        res = self.client.patch(f"/api/files/{self.f.id}/metadata/", {"review_notes": "Good. Add: she is the sponsor."}, format="json").json()
        self.assertEqual((res["meta_state"], res["review_notes"]), (before, "Good. Add: she is the sponsor."))

    def test_validation_errors(self):
        url = f"/api/files/{self.f.id}/metadata/"
        self.assertEqual(self.client.patch(url, {"doc_type": "bogus"}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(url, {"document_date": "31/12/2026"}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(url, {"signal": "x" * 401}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(url, {"unrelated": 1}, format="json").status_code, 400)
        self.f.refresh_from_db()
        self.assertEqual(self.f.doc_type, "proposal")

    def test_edit_does_not_touch_last_ingested_at(self):
        before = IngestedFile.objects.get(pk=self.f.pk).last_ingested_at
        self.client.patch(f"/api/files/{self.f.id}/metadata/", {"signal": "changed"}, format="json")
        self.assertEqual(IngestedFile.objects.get(pk=self.f.pk).last_ingested_at, before)

    def test_confirm_action(self):
        url = f"/api/files/{self.f.id}/confirm/"
        self.assertEqual(self.client.post(url).json()["meta_state"], "confirmed")
        self.client.patch(f"/api/files/{self.f.id}/metadata/", {"signal": "edited"}, format="json")
        self.assertEqual(self.client.post(url).json()["meta_state"], "corrected")  # corrected stays corrected
        bad = make_file(self.topic, "bad.docx", "", extraction_error="boom", meta_state="error")
        self.assertEqual(self.client.post(f"/api/files/{bad.id}/confirm/").status_code, 400)
        blank = make_file(self.topic, "old.txt", "x")
        self.assertEqual(self.client.post(f"/api/files/{blank.id}/confirm/").json()["meta_state"], "confirmed")

    def test_serializers_expose_metadata_and_keep_old_keys(self):
        detail = self.client.get(f"/api/topics/{self.topic.id}/").json()
        row = detail["files"][0]
        for key in ("id", "path", "display_name", "source_method", "artifact_type", "content_hash", "extraction_error",
                    "first_added_at", "last_ingested_at", "meta"):
            self.assertIn(key, row)
        for key in ("doc_type", "doc_type_label", "doc_type_source", "topic_tags", "topic_tag_list", "signal",
                    "meta_state", "review_notes", "version_label", "document_date", "meta_seconds"):
            self.assertIn(key, row)
        for key in ("id", "name", "topic_type", "one_liner", "related_people", "created_at", "auto_created", "folders",
                    "file_count", "last_activity_at", "latest_risks_and_gaps", "files", "messages"):
            self.assertIn(key, detail)
        src = self.client.get("/api/sources/").json()[0]
        for key in ("id", "topic_id", "topic_name", "topic_type", "path", "display_name", "source_method", "artifact_type", "first_added_at"):
            self.assertIn(key, src)
        self.assertEqual(src["doc_type"], "proposal")

    def test_other_existing_endpoint_shapes(self):
        t = self.topic
        ExtractedMeta.objects.create(ingested_file=self.f, meeting_title="M", meeting_date=date(2026, 8, 1))
        Task.objects.create(topic=t, title="Do", due_date=date(2026, 8, 2))
        CaseSummary.objects.create(topic=t, current_state="x")
        cal = self.client.get("/api/calendar/").json()
        meeting = next(e for e in cal if e["kind"] == "meeting")
        for key in ("kind", "date", "topic_id", "topic_name", "topic_type", "id", "file_display_name", "file_path",
                    "meeting_title", "meeting_date", "meeting_time", "attendees", "agenda", "achieved"):
            self.assertIn(key, meeting)
        stats = self.client.get("/api/stats/").json()
        for key in ("topics", "files", "meetings", "tasks_backlog", "tasks_in_progress", "tasks_done"):
            self.assertIn(key, stats)
        self.assertEqual(self.client.get(f"/api/topics/{t.id}/latest_summary/").status_code, 200)
        self.assertEqual(self.client.get("/api/search/?q=Northwind").status_code, 200)


class MetadataCsvTests(TestCase):
    def setUp(self):
        self.client = api_client()
        self.topic = Topic.objects.create(name="Northwind", topic_type="opportunity")
        self.p = make_file(self.topic, "Northwind_Proposal_V1.docx", "x", path="/proj/proposals/Northwind_Proposal_V1.docx",
                           doc_type="proposal", topic_tags="Solution, Commercial", signal="V1 proposes a pilot", meta_state="generated_unconfirmed",
                           meta_seconds=27.54, version_label="V1", document_date=date(2026, 8, 28), review_notes="check price")
        self.e = make_file(self.topic, "intro.eml", "From: Priya Menon <p@x.com>\nTo: Venki <v@m.com>\nDate: Mon, 03 Aug 2026 10:00:00 +0530\nSubject: Intro\n\nHi",
                           doc_type="email", signal='=HYPERLINK("http://evil")', meta_state="confirmed", document_date=date(2026, 8, 3))
        self.err = make_file(self.topic, "plan.xlsx", "", doc_type="project_plan", meta_state="error", extraction_error="Could not read the sheet")
        self.topic.folders.create(path="/proj")

    def rows(self, res):
        self.assertEqual(res.status_code, 200)
        self.assertIn("text/csv", res["Content-Type"])
        return list(csv.reader(io.StringIO(res.content.decode("utf-8"))))

    def test_columns_exact_and_in_order(self):
        rows = self.rows(self.client.get(f"/api/topics/{self.topic.id}/metadata_csv/"))
        self.assertEqual(
            rows[0],
            ["file", "document_type", "topic_tags", "signal", "state", "seconds", "review_notes",
             "phase", "version", "doc_date", "participants"],
        )
        self.assertEqual(rows[0][:7], ["file", "document_type", "topic_tags", "signal", "state", "seconds", "review_notes"])

    def test_values_labels_and_sorting(self):
        rows = self.rows(self.client.get(f"/api/topics/{self.topic.id}/metadata_csv/"))
        body = {r[0]: r for r in rows[1:]}
        prop = body["proposals/Northwind_Proposal_V1.docx"]  # relative to the topic folder
        self.assertEqual(prop[1], "Proposal")
        self.assertEqual((prop[2], prop[4], prop[5], prop[6]), ("Solution, Commercial", "generated_unconfirmed", "27.5", "check price"))
        self.assertEqual((prop[7], prop[8], prop[9]), ("Pre-sales", "V1", "2026-08-28"))
        mail = body["intro.eml"]
        self.assertEqual(mail[1], "Email")
        self.assertEqual(mail[10], "Priya Menon; Venki")
        self.assertTrue(mail[3].startswith("'="))  # formula injection neutralised
        err = body["plan.xlsx"]
        self.assertEqual((err[1], err[4], err[6]), ("Project Plan", "error", "Could not read the sheet"))
        self.assertEqual(err[7], "Delivery")
        dates = [r[9] or "9999" for r in rows[1:]]
        self.assertEqual(dates, sorted(dates))

    def test_all_document_type_labels(self):
        self.assertEqual(doc_meta.DOC_TYPE_LABELS["mom_action_items"], "MoM / Action Items")
        self.assertEqual(doc_meta.DOC_TYPE_LABELS["prd_brd"], "PRD / BRD")
        self.assertEqual(doc_meta.DOC_TYPE_LABELS["meeting_transcript"], "Meeting Transcript")

    def test_global_variant_adds_topic_column(self):
        other = Topic.objects.create(name="Other")
        make_file(other, "o.txt", "x", doc_type="notes")
        rows = self.rows(self.client.get("/api/metadata_csv/"))
        self.assertEqual(rows[0][:11], ["file", "document_type", "topic_tags", "signal", "state", "seconds", "review_notes", "phase", "version", "doc_date", "participants"])
        self.assertEqual(rows[0][11], "topic")
        self.assertEqual({r[11] for r in rows[1:]}, {"Northwind", "Other"})
        only = self.rows(self.client.get(f"/api/metadata_csv/?topic_id={other.id}"))
        self.assertEqual(len(only), 2)

    def test_non_sales_topic_has_blank_phase(self):
        t = Topic.objects.create(name="Eng", topic_type="project")
        make_file(t, "a.txt", "x", doc_type="notes")
        rows = self.rows(self.client.get(f"/api/topics/{t.id}/metadata_csv/"))
        self.assertEqual(rows[1][7], "")


# ---------------------------------------------------------------------------
# F. Speaker turns (parser + summary) and contacts
# ---------------------------------------------------------------------------
class SpeakerTurnTests(TestCase):
    def test_teams_parse(self):
        turns = extraction.parse_speaker_turns(TEAMS_TEXT)
        self.assertEqual([t["speaker"] for t in turns], ["Siddharth Subramanian", "Ayush Sinha A", "Vijay Krishna Ramachandran"])
        self.assertEqual([t["ts"] for t in turns], ["0:29", "1:13", "10:16:03"])
        self.assertIn("exactly the shape that I want", turns[0]["text"])
        self.assertIn("edit that also", turns[0]["text"])
        # the stray "Okay." glued before the next name belongs to the previous speaker
        self.assertTrue(turns[1]["text"].rstrip().endswith("Okay."))
        # the system line is not speech
        self.assertNotIn("stopped transcription", turns[2]["text"])
        self.assertTrue(set(turns[0]) == {"speaker", "ts", "text"})

    def test_whisper_text_is_unattributed(self):
        self.assertEqual(extraction.parse_speaker_turns(WHISPER_TEXT), [])
        self.assertEqual(extraction.parse_speaker_turns(""), [])
        self.assertEqual(extraction.parse_speaker_turns("Just a paragraph of prose with no speakers."), [])

    def test_summary_counts_and_first_three_substantive_lines(self):
        turns = [
            {"speaker": "A", "ts": "0:01", "text": "Yeah. Sure. We will ship the first timeline version on Friday. The demo needs a proper recap page. Also fix the login."},
            {"speaker": "B", "ts": "0:30", "text": "Okay."},
            {"speaker": "A", "ts": "1:00", "text": "One more substantive sentence about the quotation approval process."},
        ]
        summary = extraction.summarize_speakers(turns)
        a = summary[0]
        self.assertEqual((a["speaker"], a["turns"]), ("A", 2))
        self.assertEqual(len(a["first_lines"]), 3)
        self.assertTrue(a["first_lines"][0].startswith("We will ship"))
        b = summary[1]
        self.assertEqual((b["speaker"], b["turns"], b["words"], b["first_lines"]), ("B", 1, 1, []))

    def test_real_transcripts_folder_parses_if_present(self):
        folder = r"D:\AI Workspace\Maarga OrientMe\Transcripts"
        if not os.path.isdir(folder):
            self.skipTest("Transcripts folder not available")
        parsed = 0
        for name in os.listdir(folder):
            if name.endswith(".docx"):
                turns = extraction.parse_speaker_turns(ingestion.extract_text(os.path.join(folder, name)))
                self.assertTrue(turns, name)
                self.assertNotIn("stopped transcription", " ".join(t["text"] for t in turns))
                parsed += 1
        self.assertGreater(parsed, 0)


class ContactTests(TestCase):
    def setUp(self):
        self.client = api_client()
        self.topic = Topic.objects.create(name="N", related_people="Priya Menon")
        ExtractedMeta.objects.create(
            ingested_file=make_file(self.topic, "a", "x"),
            attendees=["Not captured (auto-transcribed; no speaker diarization)", "Rahul Iyer"])

    def test_suggestions_never_auto_create(self):
        res = self.client.get(f"/api/topics/{self.topic.id}/contacts/").json()
        self.assertEqual(res["contacts"], [])
        self.assertEqual(sorted(res["suggestions"]), ["Priya Menon", "Rahul Iyer"])
        self.assertEqual(Contact.objects.count(), 0)

    def test_crud_and_seed(self):
        url = f"/api/topics/{self.topic.id}/contacts/"
        c = self.client.post(url, {"name": "Rahul Iyer", "organisation": "Northwind", "role": "IT head", "side": "customer"}, format="json")
        self.assertEqual(c.status_code, 201)
        self.assertEqual(self.client.post(url, {"name": "rahul iyer"}, format="json").status_code, 400)
        self.assertEqual(self.client.post(url, {"name": "X", "side": "weird"}, format="json").status_code, 400)
        seeded = self.client.post(url + "seed/").json()
        self.assertEqual(seeded["created"], 1)  # only Priya; Rahul exists
        patched = self.client.patch(f"{url}{c.json()['id']}/", {"role": "CIO"}, format="json").json()
        self.assertEqual(patched["role"], "CIO")
        self.assertEqual(self.client.delete(f"{url}{c.json()['id']}/").status_code, 204)


# ---------------------------------------------------------------------------
# G. Bug fixes
# ---------------------------------------------------------------------------
class AttendeeFixTests(TestCase):
    def test_not_captured_is_never_an_attendee(self):
        meta = extraction.extract_meeting_metadata(WHISPER_TEXT)
        self.assertEqual(meta["attendees"], [])
        self.assertEqual(extraction.extract_meeting_metadata("Attendees: Not captured (engine did not report speakers)")["attendees"], [])

    def test_split_on_comma_and_semicolon(self):
        meta = extraction.extract_meeting_metadata("Meeting: Sync\nAttendees: Priya Menon; Rahul Iyer, Venki ; Gokul")
        self.assertEqual(meta["attendees"], ["Priya Menon", "Rahul Iyer", "Venki", "Gokul"])

    def test_real_names_with_not_in_them_survive(self):
        meta = extraction.extract_meeting_metadata("Attendees: Nott Captain, Anna Capture")
        self.assertEqual(meta["attendees"], ["Nott Captain", "Anna Capture"])

    def test_clean_attendees_read_filter(self):
        self.assertEqual(extraction.clean_attendees(["Not captured (x)", "Ayush", "", None, "N/A"]), ["Ayush"])

    def test_still_counts_as_meeting_for_whisper_header(self):
        self.assertTrue(extraction.looks_like_meeting_transcript(WHISPER_TEXT))


class DocxTableTests(TestCase):
    def build(self, with_table=True):
        import docx

        d = docx.Document()
        d.add_paragraph("Quotation for Northwind")
        d.add_paragraph("")
        d.add_paragraph("Scope below.")
        if with_table:
            t = d.add_table(rows=3, cols=3)
            rows = [("Item", "Qty", "Price"), ("Pilot", "1", "50,000"), ("Support", "6", "1,200")]
            for r, vals in enumerate(rows):
                for c, v in enumerate(vals):
                    t.cell(r, c).text = v
            d.add_paragraph("Valid until 30 Sep 2026.")
        path = os.path.join(tempfile.mkdtemp(), "q.docx")
        d.save(path)
        return path

    def test_table_cells_are_read_in_document_order(self):
        text = ingestion.extract_text(self.build())
        lines = text.split("\n")
        self.assertEqual(lines[0], "Quotation for Northwind")
        self.assertIn("Item\tQty\tPrice", lines)
        self.assertIn("Pilot\t1\t50,000", lines)
        self.assertLess(lines.index("Scope below."), lines.index("Pilot\t1\t50,000"))
        self.assertLess(lines.index("Support\t6\t1,200"), lines.index("Valid until 30 Sep 2026."))

    def test_paragraph_only_docx_is_unchanged(self):
        import docx

        path = self.build(with_table=False)
        expected = "\n".join(p.text for p in docx.Document(path).paragraphs if p.text.strip())
        self.assertEqual(ingestion.extract_text(path), expected)


class HtmlEmailTests(TestCase):
    def write_eml(self, html=None, plain=None):
        from email.message import EmailMessage

        msg = EmailMessage()
        msg["From"], msg["To"], msg["Subject"], msg["Date"] = "a@b.com", "c@d.com", "Hello", "Mon, 03 Aug 2026 10:00:00 +0530"
        if plain is not None and html is not None:
            msg.set_content(plain)
            msg.add_alternative(html, subtype="html")
        elif html is not None:
            msg.set_content(html, subtype="html")
        else:
            msg.set_content(plain)
        path = os.path.join(tempfile.mkdtemp(), "m.eml")
        with open(path, "wb") as fh:
            fh.write(bytes(msg))
        return path

    def test_html_only_body_is_stripped(self):
        html = "<html><head><style>p{color:red}</style><title>x</title></head><body><p>Hi <b>Venki</b>,</p><p>Shall we meet &amp; talk on Thursday?</p><script>alert(1)</script><table><tr><td>A</td><td>B</td></tr></table></body></html>"
        text = ingestion.extract_text(self.write_eml(html=html))
        self.assertIn("Subject: Hello", text)
        self.assertIn("Hi Venki,", text)
        self.assertIn("Shall we meet & talk on Thursday?", text)
        self.assertNotIn("<", text)
        self.assertNotIn("color:red", text)
        self.assertNotIn("alert(1)", text)
        self.assertIn("A\tB", text)

    def test_plain_part_is_preferred_and_untouched(self):
        text = ingestion.extract_text(self.write_eml(plain="Plain <body> text", html="<p>HTML version</p>"))
        self.assertIn("Plain <body> text", text)
        self.assertNotIn("HTML version", text)

    def test_html_to_text_handles_garbage(self):
        self.assertEqual(ingestion.html_to_text("no tags"), "no tags")
        self.assertEqual(ingestion.html_to_text(""), "")


class UploadDocumentTests(TestCase):
    """G5: a plain document uploaded through the Transcripts-style upload is not
    renamed to "<name> - Transcript.txt" and gets no recording event; a real
    transcript still takes the recording path."""

    def setUp(self):
        self.client = api_client()
        self.topic = Topic.objects.create(name="Northwind", topic_type="opportunity")
        self.media = tempfile.mkdtemp()

    def upload(self, name, content):
        from django.core.files.uploadedfile import SimpleUploadedFile

        with override_settings(MEDIA_ROOT=self.media):
            return self.client.post(
                "/api/recording_events/upload/",
                {"topic_id": self.topic.id, "files": SimpleUploadedFile(name, content)},
                format="multipart",
            )

    def test_document_keeps_its_name_and_is_not_a_recording(self):
        res = self.upload("Northwind_Proposal_V1.txt", b"Northwind Proposal\n\nA six week pilot is proposed for warehouse exception handling.")
        self.assertEqual(res.status_code, 200, res.content)
        result = res.json()["results"][0]
        self.assertTrue(result["document"])
        self.assertEqual(RecordingEvent.objects.count(), 0)
        f = self.topic.files.get()
        self.assertEqual(f.display_name, "Northwind_Proposal_V1.txt")
        self.assertNotIn("Transcript", f.display_name)
        self.assertEqual((f.doc_type, f.version_label), ("proposal", "V1"))
        self.assertTrue(os.path.exists(f.path))
        self.assertFalse(any("Transcript" in n for n in os.listdir(os.path.dirname(f.path))))

    def test_transcript_text_still_takes_the_recording_path(self):
        res = self.upload("Weekly Meeting-20261005_044517UTC-Meeting Recording.txt", TEAMS_TEXT.encode("utf-8"))
        self.assertEqual(res.status_code, 200, res.content)
        event = RecordingEvent.objects.get()
        self.assertEqual(event.status, "routed")
        self.assertTrue(event.ingested_file.display_name.endswith(" - Transcript.txt"))
        self.assertEqual(len(event.speaker_turns), 3)  # Teams names parsed in code right away

    def test_email_attachment_is_a_document(self):
        eml = b"From: a@b.com\nTo: c@d.com\nSubject: Intro\nDate: Mon, 03 Aug 2026 10:00:00 +0530\n\nHello there, can we meet on Thursday about the pilot?\nAttendees: Priya, Rahul\n"
        res = self.upload("intro.eml", eml)
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(RecordingEvent.objects.count(), 0)
        self.assertEqual(self.topic.files.get().doc_type, "email")

    def test_ingest_into_topic_names_documents_by_themselves(self):
        path = os.path.join(self.media, "Quote_v1.txt")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("Quotation\n\nFixed fee of the price for the pilot.")
        obj = recordings._ingest_into_topic(self.topic, path, "Quotation\n\nFixed fee of the price for the pilot.", refresh_brief=False)
        self.assertEqual((obj.path, obj.display_name), (path, "Quote_v1.txt"))


# ---------------------------------------------------------------------------
# D + H. Brief prompt
# ---------------------------------------------------------------------------
class BriefPromptTests(TestCase):
    def prompt_for(self, topic, files=None):
        captured = {}

        def fake(prompt, settings):
            captured["prompt"] = prompt
            return {"is_conversational": False, "answer_summary": "x", "current_state": "x", "sources_used": []}

        files = files if files is not None else list(topic.files.all())
        with mock.patch.object(llm_brief, "_via_openrouter", side_effect=fake):
            llm_brief.generate_brief(topic, "Give me the current state", files, settings={"llm_api_key": "test-key-not-real"})
        return captured["prompt"]

    def test_plain_topic_without_notes_gets_exactly_the_old_prompt(self):
        topic = Topic.objects.create(name="Eng", topic_type="project")
        f = make_file(topic, "a.txt", "Some text about the project.", doc_type="notes", version_label="V1", document_date=date(2026, 8, 1))
        prompt = self.prompt_for(topic, [f])
        baseline = llm_brief.PROMPT_TEMPLATE.format(
            topic_name="Eng", topic_type="Project", manual_state_block="", manual_notes_block="", sales_recap_block="",
            question="Give me the current state", files_block=f"--- FILE: {f.path} ---\nSome text about the project.")
        self.assertEqual(prompt, baseline)
        self.assertNotIn("HUMAN NOTES", prompt)
        self.assertNotIn("sales", prompt.split("Scope:")[1])
        self.assertNotIn("[Document:", prompt)
        # the section between the scope line and the question is exactly what it always was
        self.assertIn("Scope: Eng (Project)\n\nQuestion being asked:", prompt)

    def test_notes_block_is_dated_trusted_and_next_to_project_state(self):
        topic = Topic.objects.create(name="Eng", topic_type="project")
        from .models import ProjectState

        ProjectState.objects.create(topic=topic, current_position="We are live.")
        meta = ExtractedMeta.objects.create(ingested_file=make_file(topic, "m.txt", "x"), meeting_title="Discovery 1", meeting_date=date(2026, 8, 14))
        MeetingNote.objects.create(topic=topic, meta=meta, text="Sponsor is the ops head.\nBudget not discussed.", author="ayush")
        MeetingNote.objects.create(topic=topic, text="Customer wants a pilot before November.", author="venki")
        prompt = self.prompt_for(topic)
        self.assertIn("Human notes (trusted), oldest first:", prompt)
        self.assertIn("Sponsor is the ops head. Budget not discussed. - ayush", prompt)
        self.assertIn("(about the meeting 'Discovery 1', 14 Aug 2026)", prompt)
        self.assertIn("Customer wants a pilot before November. - venki", prompt)
        self.assertIn("HIGH-SIGNAL and TRUSTED", prompt)
        self.assertIn("conflicts with a source", prompt)
        self.assertLess(prompt.index("Manually confirmed current state"), prompt.index("Human notes (trusted)"))
        self.assertLess(prompt.index("Human notes (trusted)"), prompt.index("Question being asked"))
        self.assertNotIn("sales opportunity", prompt)

    def test_opportunity_and_customer_get_sales_structure(self):
        for kind in ("opportunity", "customer"):
            topic = Topic.objects.create(name=f"N-{kind}", topic_type=kind, stage="discovery", customer_name="Northwind Retail")
            prop = make_file(topic, "Northwind_Proposal_V2.docx", "Proposal text here.", doc_type="proposal", version_label="V2", document_date=date(2026, 9, 9))
            meta_file = make_file(topic, "d2.txt", "Transcript text.", doc_type="meeting_transcript")
            ExtractedMeta.objects.create(ingested_file=meta_file, meeting_title="Discovery 2", meeting_date=date(2026, 8, 21), meeting_type="discovery", meeting_sequence=2)
            MeetingNote.objects.create(topic=topic, text="Pricing is the blocker.", author="a")
            prompt = self.prompt_for(topic)
            self.assertIn(f"This topic is a sales {kind}.", prompt)
            self.assertIn("Customer: Northwind Retail.", prompt)
            self.assertIn("Stage (set by a person): Discovery.", prompt)
            for needle in ("Account snapshot", "Stakeholders with role and stance", "DATE ORDER", "VERSION (V1, V2, Final)",
                           "Open items", "Next step", "current_state", "customer_thoughts", "history", "proposed_solution",
                           "promises_made", "risks_and_gaps", "next_step"):
                self.assertIn(needle, prompt)
            self.assertIn("Pricing is the blocker.", prompt)
            self.assertIn(f"--- FILE: {prop.path} ---\n[Document: Proposal, version V2, dated 2026-09-09]\nProposal text here.", prompt)
            self.assertIn("[Document: Meeting Transcript, Discovery 2, dated 2026-08-21]", prompt)
            # the original instructions are all still there
            self.assertIn("You are OrientMe, an orientation assistant", prompt)
            self.assertIn("Call submit_orientme_brief with the finished brief.", prompt)

    def test_other_topic_types_do_not_get_the_sales_block(self):
        for kind in ("project", "area", "process", "person"):
            topic = Topic.objects.create(name=f"T-{kind}", topic_type=kind)
            self.assertNotIn("sales", self.prompt_for(topic).split("Scope:")[1].split("Question being asked")[0])

    def test_call_structured_requires_a_key(self):
        with self.assertRaises(llm_brief.MissingApiKeyError):
            llm_brief.call_structured("p", "t", "d", {"type": "object"}, {})


# ---------------------------------------------------------------------------
# I. Who spoke what (Transcripts detail page)
# ---------------------------------------------------------------------------
def lines_of(n, prefix="Line"):
    return [(f"{i // 60:02d}:{i % 60:02d}", f"{prefix} {i} of the recording says something quite specific about topic {i}.") for i in range(1, n + 1)]


def whisper_event(n=6, text=None, **kw):
    topic = kw.pop("topic", None) or Topic.objects.create(name="W")
    body = text or "Title: T\nDate: 2026-10-05\nAttendees: Not captured (auto-transcribed; no speaker diarization)\n\n" + "\n".join(
        f"[{ts}] {t}" for ts, t in lines_of(n))
    f = make_file(topic, "w.txt", body)
    return RecordingEvent.objects.create(file_path=f"/rec/{f.id}.mp4", topic=topic, ingested_file=f, status="routed", **kw)


class SpeakerChunkingTests(TestCase):
    def test_chunk_ranges_overlap_and_cover_everything(self):
        lines = lines_of(200)
        ranges = recordings._speaker_chunk_ranges(lines, max_chars=1500, overlap=10)
        self.assertGreater(len(ranges), 2)
        covered = []
        for start, end, label_from in ranges:
            self.assertLess(label_from, end)
            self.assertLessEqual(start, label_from)
            self.assertLessEqual(label_from - start, 10)  # at most 10 context lines
            covered.extend(range(label_from, end))
        self.assertEqual(covered, list(range(200)))  # every line is "to label" exactly once
        self.assertEqual(ranges[0][0], 0)
        self.assertEqual(ranges[-1][1], 200)

    def test_single_short_transcript_is_one_chunk(self):
        self.assertEqual(recordings._speaker_chunk_ranges(lines_of(5)), [(0, 5, 0)])

    def test_huge_single_line_still_progresses(self):
        lines = [("00:01", "x" * 20000), ("00:02", "short")]
        ranges = recordings._speaker_chunk_ranges(lines, max_chars=8000, overlap=10)
        self.assertEqual(sum(e - lf for _s, e, lf in ranges), 2)


class SpeakerTranscriptTests(TestCase):
    key = {"llm_api_key": "test-key-not-real", "llm_model": "m", "llm_base_url": "http://x"}

    def run_generate(self, event, call, settings=None):
        return recordings.generate_speaker_transcript(event, settings=settings or self.key, call=call)

    def test_single_speaker_is_a_valid_result(self):
        event = whisper_event(6)
        call = mock.Mock(return_value={"groups": [{"start_line": 1, "end_line": 6, "speaker": "Speaker 1"}]})
        _e, ok, msg = self.run_generate(event, call)
        self.assertTrue(ok)
        self.assertEqual(msg, "Transcript generated.")
        event.refresh_from_db()
        self.assertEqual({t["speaker"] for t in event.speaker_turns}, {"Speaker 1"})
        self.assertEqual(event.speaker_turns_source, "llm")
        self.assertEqual(call.call_count, 1)
        # every real word and timestamp is still there, from code
        self.assertEqual(len(event.speaker_transcript.splitlines()), 6)
        self.assertEqual(event.speaker_transcript.splitlines()[0], "[00:01] Speaker 1: Line 1 of the recording says something quite specific about topic 1.")
        joined = " ".join(t["text"] for t in event.speaker_turns)
        for i in range(1, 7):
            self.assertIn(f"Line {i} of the recording", joined)
        prompt = call.call_args[0][0]
        self.assertIn("only ONE person speaks", prompt)  # the prompt allows a single speaker
        self.assertIn("never invent a second speaker", prompt)

    def test_two_speakers_become_turns_with_first_timestamps(self):
        event = whisper_event(6)
        call = mock.Mock(return_value={"groups": [
            {"start_line": 1, "end_line": 2, "speaker": "Speaker 1"},
            {"start_line": 3, "end_line": 4, "speaker": "speaker 2"},
            {"start_line": 5, "end_line": 6, "speaker": "Speaker 1"},
        ]})
        self.run_generate(event, call)
        event.refresh_from_db()
        self.assertEqual([(t["speaker"], t["ts"]) for t in event.speaker_turns],
                         [("Speaker 1", "00:01"), ("Speaker 2", "00:03"), ("Speaker 1", "00:05")])
        self.assertTrue(all(t["edited"] is False for t in event.speaker_turns))

    def test_model_never_gets_to_rewrite_text_and_bad_labels_are_rejected(self):
        event = whisper_event(4)
        call = mock.Mock(return_value={"groups": [
            {"start_line": 1, "end_line": 2, "speaker": "Assistant"},  # forbidden label -> ignored
            {"start_line": 1, "end_line": 4, "speaker": "Gokul", "text": "fabricated words"},
        ]})
        self.run_generate(event, call)
        event.refresh_from_db()
        self.assertEqual({t["speaker"] for t in event.speaker_turns}, {"Gokul"})
        self.assertNotIn("fabricated", event.speaker_transcript)
        self.assertNotIn("fabricated", str(event.speaker_turns))

    def test_long_transcript_is_chunked_with_consistent_labels(self):
        event = whisper_event(120)
        prompts = []

        def fake(prompt, tool, desc, schema, settings, timeout=120):
            prompts.append(prompt)
            # label by line number parity of the 'to label' range given in the prompt
            import re as _re

            m = _re.search(r"Label lines (\d+)-(\d+)\.", prompt)
            lo, hi = (int(m.group(1)), int(m.group(2))) if m else (1, int(_re.findall(r"^(\d+)\. ", prompt, _re.M)[-1]))
            groups = [{"start_line": i, "end_line": i, "speaker": "Speaker 1" if (i // 10) % 2 == 0 else "Speaker 2"} for i in range(lo, hi + 1)]
            return {"groups": groups}

        with mock.patch.object(recordings, "SPEAKER_CHUNK_CHARS", 2000):
            _e, ok, msg = recordings.generate_speaker_transcript(event, settings=self.key, call=fake)
        self.assertTrue(ok)
        event.refresh_from_db()
        labelled = event.speaker_transcript.splitlines()
        self.assertEqual(len(labelled), 120)
        self.assertFalse(any("Unattributed" in ln for ln in labelled))

    def test_chunk_two_gets_known_speakers_and_labelled_context(self):
        lines = lines_of(60)
        prompts = []

        def fake(prompt, tool, desc, schema, settings, timeout=120):
            import re as _re

            prompts.append(prompt)
            m = _re.search(r"Label lines (\d+)-(\d+)\.", prompt)
            lo = int(m.group(1)) if m else 1
            hi = int(m.group(2)) if m else int(_re.findall(r"^(\d+)\. ", prompt, _re.M)[-1])
            return {"groups": [{"start_line": lo, "end_line": hi, "speaker": "Gokul" if lo == 1 else "gokul"}]}

        with mock.patch.object(recordings, "_speaker_chunk_ranges", lambda ln: [(0, 30, 0), (20, 60, 30)]):
            labels, unattributed, error = recordings._label_lines_with_llm(lines, self.key, call=fake)
        self.assertEqual((unattributed, error), (0, None))
        self.assertEqual(len(prompts), 2)
        self.assertNotIn("earlier parts", prompts[0])
        self.assertIn("Speakers identified in earlier parts of this meeting: Gokul.", prompts[1])
        self.assertIn("Lines 21-30 are context", prompts[1])
        self.assertIn("Label lines 31-60.", prompts[1])
        self.assertIn("21. [Gokul] Line 21", prompts[1])
        self.assertIn("31. Line 31", prompts[1])
        self.assertEqual(set(labels.values()), {"Gokul"})  # 'gokul' unified with the known label

    def test_unlabelled_lines_are_unattributed_not_inherited(self):
        event = whisper_event(6)
        # the model only ever covers lines 1-3 (even after the retry)
        call = mock.Mock(return_value={"groups": [{"start_line": 1, "end_line": 3, "speaker": "Speaker 1"}]})
        _e, ok, msg = self.run_generate(event, call)
        self.assertTrue(ok)
        self.assertEqual(call.call_count, 2)  # one retry for the gap
        self.assertIn("3 line(s) could not be attributed", msg)
        event.refresh_from_db()
        self.assertEqual([t["speaker"] for t in event.speaker_turns], ["Speaker 1", "Unattributed"])

    def test_gap_is_filled_by_the_retry(self):
        event = whisper_event(6)
        call = mock.Mock(side_effect=[
            {"groups": [{"start_line": 1, "end_line": 3, "speaker": "Speaker 1"}]},
            {"groups": [{"start_line": 4, "end_line": 6, "speaker": "Speaker 2"}]},
        ])
        _e, ok, msg = self.run_generate(event, call)
        self.assertEqual(msg, "Transcript generated.")
        event.refresh_from_db()
        self.assertEqual([t["speaker"] for t in event.speaker_turns], ["Speaker 1", "Speaker 2"])

    def test_no_api_key_degrades_to_one_placeholder_speaker_and_mentions_the_key(self):
        event = whisper_event(4)
        call = mock.Mock()
        _e, ok, msg = recordings.generate_speaker_transcript(event, settings={"llm_api_key": ""}, call=call)
        self.assertTrue(ok)
        call.assert_not_called()
        self.assertIn("API key", msg)
        self.assertNotIn("Ollama", msg)
        event.refresh_from_db()
        self.assertEqual({t["speaker"] for t in event.speaker_turns}, {"Speaker 1"})
        self.assertEqual(event.speaker_turns_source, "code")
        self.assertEqual(len(event.speaker_transcript.splitlines()), 4)

    def test_api_failure_degrades_the_same_way(self):
        event = whisper_event(4)
        call = mock.Mock(side_effect=RuntimeError("boom with secret-looking-text"))
        _e, ok, msg = self.run_generate(event, call)
        self.assertTrue(ok)
        self.assertIn("API key", msg)
        self.assertNotIn("secret-looking-text", msg)
        event.refresh_from_db()
        self.assertEqual({t["speaker"] for t in event.speaker_turns}, {"Speaker 1"})
        self.assertEqual(event.speaker_turns_source, "code")

    def test_partial_failure_marks_only_the_failed_chunk(self):
        event = whisper_event(40)
        state = {"n": 0}

        def fake(prompt, tool, desc, schema, settings, timeout=120):
            import re as _re

            state["n"] += 1
            m = _re.search(r"Label lines (\d+)-(\d+)\.", prompt)
            if m:
                raise RuntimeError("second chunk always fails")
            hi = int(_re.findall(r"^(\d+)\. ", prompt, _re.M)[-1])
            return {"groups": [{"start_line": 1, "end_line": hi, "speaker": "Speaker 1"}]}

        with mock.patch.object(recordings, "_speaker_chunk_ranges", lambda ln: [(0, 20, 0), (10, 40, 20)]):
            _e, ok, msg = self.run_generate(event, fake)
        self.assertIn("could not be attributed", msg)
        event.refresh_from_db()
        self.assertEqual([t["speaker"] for t in event.speaker_turns][-1], "Unattributed")

    def test_teams_format_is_parsed_in_code_without_any_llm_call(self):
        topic = Topic.objects.create(name="T")
        f = make_file(topic, "t.docx", TEAMS_TEXT)
        event = RecordingEvent.objects.create(file_path="/rec/t.docx", topic=topic, ingested_file=f, status="routed")
        call = mock.Mock()
        _e, ok, msg = recordings.generate_speaker_transcript(event, settings=self.key, call=call)
        call.assert_not_called()
        self.assertTrue(ok)
        self.assertIn("no LLM was used", msg)
        event.refresh_from_db()
        self.assertEqual([t["speaker"] for t in event.speaker_turns], ["Siddharth Subramanian", "Ayush Sinha A", "Vijay Krishna Ramachandran"])
        self.assertEqual(event.speaker_turns_source, "teams")
        self.assertTrue(event.speaker_transcript.startswith("[0:29] Siddharth Subramanian: And this is exactly"))

    def test_no_timestamped_lines_and_empty_text(self):
        topic = Topic.objects.create(name="T")
        f = make_file(topic, "p.txt", "Just prose with no timestamps at all.")
        event = RecordingEvent.objects.create(file_path="/rec/p", topic=topic, ingested_file=f, status="routed")
        self.assertFalse(recordings.generate_speaker_transcript(event, settings=self.key, call=mock.Mock())[1])
        empty = RecordingEvent.objects.create(file_path="/rec/e", status="needs_review")
        self.assertFalse(recordings.generate_speaker_transcript(empty, settings=self.key, call=mock.Mock())[1])

    def test_local_model_is_never_called(self):
        event = whisper_event(4)
        with mock.patch.object(recordings.local_llm, "call_local_llm") as local:
            recordings.generate_speaker_transcript(event, settings=self.key, call=mock.Mock(return_value={"groups": []}))
            recordings.generate_speaker_transcript(event, settings={"llm_api_key": ""})
        local.assert_not_called()


class SpeakerEndpointTests(TestCase):
    def setUp(self):
        self.client = api_client()

    def test_rename_speaker_marks_edited_and_updates_text(self):
        event = whisper_event(4)
        recordings.generate_speaker_transcript(
            event, settings={"llm_api_key": "test-key-not-real"},
            call=mock.Mock(return_value={"groups": [{"start_line": 1, "end_line": 2, "speaker": "Speaker 1"}, {"start_line": 3, "end_line": 4, "speaker": "Speaker 2"}]}))
        res = self.client.post(f"/api/recording_events/{event.id}/rename_speaker/", {"old": "Speaker 1", "new": "  Gokul "}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(body["renamed_turns"], 1)
        self.assertEqual([(t["speaker"], t["edited"]) for t in body["speaker_turns"]], [("Gokul", True), ("Speaker 2", False)])
        self.assertTrue(body["speaker_transcript"].startswith("[00:01] Gokul: Line 1"))
        self.assertIn("] Speaker 2: Line 3", body["speaker_transcript"])
        self.assertEqual(self.client.post(f"/api/recording_events/{event.id}/rename_speaker/", {"old": "Nobody", "new": "X"}, format="json").status_code, 404)
        self.assertEqual(self.client.post(f"/api/recording_events/{event.id}/rename_speaker/", {"old": "", "new": "X"}, format="json").status_code, 400)
        self.assertEqual(self.client.post(f"/api/recording_events/{event.id}/rename_speaker/", {"old": "Gokul", "new": "x" * 61}, format="json").status_code, 400)
        event.refresh_from_db()
        self.assertEqual(event.speaker_turns[0]["speaker"], "Gokul")  # persisted

    def test_serializer_fields_and_old_keys_and_no_key_leak(self):
        Settings.set("llm_api_key", "test-key-not-real")
        event = whisper_event(3)
        with mock.patch.object(recordings.threading, "Thread"):  # a real background thread would fight the test DB
            detail = self.client.get(f"/api/recording_events/{event.id}/").json()
        recordings._auto_inflight.discard(event.id)
        for key in ("id", "file_path", "file_name", "detected_at", "recording_date", "source", "topic", "topic_name",
                    "ingested_file", "status", "engine_used", "detail", "suggested_topic_name", "suggested_topic_one_liner",
                    "suggested_topic_type", "meeting_summary", "summary_generated_at", "diarization_result",
                    "diarization_generated_at", "speaker_transcript", "speaker_transcript_generated_at"):
            self.assertIn(key, detail)
        for key in ("speaker_turns", "speaker_turns_source", "speaker_turn_count", "speaker_turns_generating", "llm_key_configured"):
            self.assertIn(key, detail)
        self.assertTrue(detail["llm_key_configured"])
        self.assertNotIn("test-key-not-real", str(detail))
        self.assertNotIn("test-key-not-real", str(self.client.get("/api/recording_events/").json()))
        listing = self.client.get("/api/recording_events/").json()
        row = listing[0] if isinstance(listing, list) else listing["results"][0]
        self.assertNotIn("speaker_turns", row)  # the list stays light
        self.assertIn("speaker_turn_count", row)

    def test_retrieve_prepares_teams_turns_inline_and_caches(self):
        topic = Topic.objects.create(name="T")
        f = make_file(topic, "t.docx", TEAMS_TEXT)
        event = RecordingEvent.objects.create(file_path="/rec/t.docx", topic=topic, ingested_file=f, status="routed")
        body = self.client.get(f"/api/recording_events/{event.id}/").json()
        self.assertEqual(len(body["speaker_turns"]), 3)
        self.assertEqual(body["speaker_turns_source"], "teams")

    def test_retrieve_uses_existing_speaker_transcript_without_calling_the_llm(self):
        Settings.set("llm_api_key", "test-key-not-real")
        event = whisper_event(
            3, speaker_transcript="[00:01] Speaker 1: Line one.\n[00:02] Speaker 1: Line two.\n[00:03] Rinky: Line three.")
        with mock.patch.object(llm_brief, "call_structured") as call:
            body = self.client.get(f"/api/recording_events/{event.id}/").json()
        call.assert_not_called()
        self.assertEqual([t["speaker"] for t in body["speaker_turns"]], ["Speaker 1", "Rinky"])
        self.assertEqual(body["speaker_turns_source"], "legacy")

    def test_retrieve_starts_background_generation_once_when_a_key_exists(self):
        Settings.set("llm_api_key", "test-key-not-real")
        event = whisper_event(4)
        with mock.patch.object(recordings.threading, "Thread") as thread:
            body = self.client.get(f"/api/recording_events/{event.id}/").json()
            self.assertEqual(thread.call_count, 1)
            thread.return_value.start.assert_called_once()
            self.assertTrue(body["speaker_turns_generating"])
            recordings._auto_inflight.discard(event.id)
            self.client.get(f"/api/recording_events/{event.id}/")
            self.assertEqual(thread.call_count, 1)  # attempted once, cached: not started again
        event.refresh_from_db()
        self.assertIsNotNone(event.speaker_turns_attempted_at)

    def test_ensure_without_key_does_nothing(self):
        event = whisper_event(4)
        with mock.patch.object(recordings.threading, "Thread") as thread:
            self.assertEqual(recordings.ensure_speaker_turns(event), "skipped")
        thread.assert_not_called()

    def test_ensure_runs_generation_and_never_raises(self):
        Settings.set("llm_api_key", "test-key-not-real")
        event = whisper_event(4)
        with mock.patch.object(llm_brief, "call_structured", return_value={"groups": [{"start_line": 1, "end_line": 4, "speaker": "Speaker 1"}]}):
            self.assertEqual(recordings.ensure_speaker_turns(event, background=False), "ready")
        event.refresh_from_db()
        self.assertEqual(len(event.speaker_turns), 1)
        self.assertEqual(recordings.ensure_speaker_turns(event), "ready")
        broken = whisper_event(4)
        with mock.patch.object(llm_brief, "call_structured", side_effect=RuntimeError("x")):
            self.assertEqual(recordings.ensure_speaker_turns(broken, background=False), "ready")  # degraded, no exception

    def test_process_new_recording_failure_in_auto_turns_does_not_break_ingestion(self):
        topic = Topic.objects.create(name="T")
        d = tempfile.mkdtemp()
        path = os.path.join(d, "Meeting-20261005_044517UTC-Meeting Recording - Transcript.txt")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(TEAMS_TEXT)
        with mock.patch.object(recordings, "ensure_speaker_turns", side_effect=RuntimeError("boom")):
            recordings.process_new_recording(path, target_topic=topic, refresh_brief=False)
        self.assertEqual(RecordingEvent.objects.get(file_path=path).status, "routed")

    def test_existing_generate_transcript_endpoint_still_works(self):
        event = whisper_event(3)
        res = self.client.post(f"/api/recording_events/{event.id}/generate_transcript/")
        self.assertEqual(res.status_code, 200)
        body = res.json()
        self.assertTrue(body["succeeded"])
        self.assertIn("API key", body["message"])
        self.assertEqual(len(body["speaker_transcript"].splitlines()), 3)
        self.assertEqual(self.client.post(f"/api/recording_events/{event.id}/generate_summary/").status_code, 200)


# ---------------------------------------------------------------------------
# Backfill command + migration
# ---------------------------------------------------------------------------
class BackfillCommandTests(TestCase):
    def test_idempotent_and_leaves_manual_values_alone(self):
        from django.core.management import call_command

        topic = Topic.objects.create(name="T")
        a = make_file(topic, "Northwind_Proposal_V1.docx", "Proposal\n\nThis proposal describes a six week pilot for the warehouse team.")
        b = make_file(topic, "Report.pdf", "x", doc_type="notes", doc_type_source="manual", meta_state="corrected", signal="mine")
        out = io.StringIO()
        call_command("backfill_doc_metadata", "--dry-run", stdout=out)
        self.assertIn("would fill metadata for 1", out.getvalue())
        a.refresh_from_db()
        self.assertEqual(a.meta_state, "")  # dry run wrote nothing
        call_command("backfill_doc_metadata", stdout=out)
        a.refresh_from_db()
        b.refresh_from_db()
        self.assertEqual((a.doc_type, a.version_label, a.meta_state), ("proposal", "V1", "generated_unconfirmed"))
        self.assertEqual((b.doc_type, b.signal, b.meta_state), ("notes", "mine", "corrected"))
        out2 = io.StringIO()
        call_command("backfill_doc_metadata", stdout=out2)
        self.assertIn("filled metadata for 0", out2.getvalue())

    def test_meetings_flag_fills_type_and_turns_without_touching_manual(self):
        from django.core.management import call_command

        topic = Topic.objects.create(name="T")
        auto = ExtractedMeta.objects.create(ingested_file=make_file(topic, "a", TEAMS_TEXT), meeting_title="Discovery 2 with X")
        manual = ExtractedMeta.objects.create(
            ingested_file=make_file(topic, "b", TEAMS_TEXT), meeting_title="Discovery 2", meeting_type="solution", meeting_type_source="manual")
        out = io.StringIO()
        call_command("backfill_doc_metadata", "--meetings", stdout=out)
        auto.refresh_from_db()
        manual.refresh_from_db()
        self.assertEqual((auto.meeting_type, auto.meeting_sequence), ("discovery", 2))
        self.assertEqual(len(auto.speaker_turns), 3)
        self.assertEqual((manual.meeting_type, manual.meeting_type_source), ("solution", "manual"))
        self.assertEqual(len(manual.speaker_turns), 3)  # turns are always derivable from the text
        out2 = io.StringIO()
        call_command("backfill_doc_metadata", "--meetings", stdout=out2)
        self.assertIn("Meetings: filled type/speaker turns for 0", out2.getvalue())


class MigrationTests(TransactionTestCase):
    """The new migration is additive: existing rows keep working with defaults."""

    migrate_from = [("core", "0026_remove_inbox")]
    migrate_to = [("core", "0027_sales_metadata_notes_speakers")]

    def test_forward_migration_keeps_rows_and_applies_defaults(self):
        executor = MigrationExecutor(connection)
        executor.migrate(self.migrate_from)
        old_apps = executor.loader.project_state(self.migrate_from).apps
        Topic = old_apps.get_model("core", "Topic")
        IngestedFile = old_apps.get_model("core", "IngestedFile")
        ExtractedMeta = old_apps.get_model("core", "ExtractedMeta")
        RecordingEvent = old_apps.get_model("core", "RecordingEvent")
        topic = Topic.objects.create(name="Old topic", topic_type="project")
        f = IngestedFile.objects.create(topic=topic, path="/old/a.txt", extracted_text="hello")
        ExtractedMeta.objects.create(ingested_file=f, meeting_title="Old meeting")
        RecordingEvent.objects.create(file_path="/old/r.mp4")

        executor = MigrationExecutor(connection)
        executor.migrate(self.migrate_to)
        new_apps = executor.loader.project_state(self.migrate_to).apps
        t = new_apps.get_model("core", "Topic").objects.get(name="Old topic")
        self.assertEqual((t.stage, t.customer_name, t.topic_type), ("", "", "project"))
        nf = new_apps.get_model("core", "IngestedFile").objects.get(path="/old/a.txt")
        self.assertEqual((nf.doc_type, nf.meta_state, nf.topic_tags, nf.signal, nf.version_label, nf.review_notes), ("", "", "", "", "", ""))
        self.assertIsNone(nf.document_date)
        self.assertIsNone(nf.meta_seconds)
        self.assertEqual(nf.extracted_text, "hello")
        nm = new_apps.get_model("core", "ExtractedMeta").objects.get(meeting_title="Old meeting")
        self.assertEqual((nm.meeting_type, nm.meeting_sequence, nm.speaker_turns), ("", None, []))
        ev = new_apps.get_model("core", "RecordingEvent").objects.get(file_path="/old/r.mp4")
        self.assertEqual((ev.speaker_turns, ev.speaker_turns_source, ev.speaker_turns_attempted_at), ([], "", None))
        self.assertEqual(new_apps.get_model("core", "MeetingNote").objects.count(), 0)
        self.assertEqual(new_apps.get_model("core", "Contact").objects.count(), 0)
